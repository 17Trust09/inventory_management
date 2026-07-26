from __future__ import annotations

import csv
import os
from datetime import datetime, timedelta
from typing import Iterable

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import ScheduledExport

from .models import InventoryItem, Overview


EXPORT_COLUMNS = [
    ("id", "ID", lambda it: it.id),
    ("name", "Name", lambda it: it.name),
    ("image", "Bild", lambda it: _item_image_url(it)),
    ("type", "Typ", lambda it: it.get_item_type_display()),
    ("quantity", "Bestand", lambda it: it.quantity),
    ("unit", "Einheit", lambda it: it.get_unit_display() if it.unit else ""),
    ("variant", "Variante", lambda it: it.variant or ""),
    ("category", "Kategorie", lambda it: it.category.name if it.category else ""),
    ("storage_location", "Lagerort", lambda it: it.storage_location.get_full_path() if it.storage_location else ""),
    ("location_letter", "Ort (Buchstabe)", lambda it: it.location_letter or ""),
    ("location_number", "Ort (Nummer)", lambda it: it.location_number or ""),
    ("location_shelf", "Ort (Fach)", lambda it: it.location_shelf or ""),
    ("min_stock", "Mindestbestand", lambda it: it.low_quantity),
    ("tags", "Tags", lambda it: ", ".join(sorted(it.application_tags.values_list("name", flat=True)))),
    ("overview", "Dashboard", lambda it: it.overview.name if it.overview else ""),
    ("maintenance_date", "Wartungsdatum", lambda it: it.maintenance_date.isoformat() if it.maintenance_date else ""),
    ("last_used", "Letzte Nutzung", lambda it: it.last_used.isoformat() if it.last_used else ""),
    ("created_at", "Erstellt am", lambda it: it.date_created.isoformat() if it.date_created else ""),
]


def get_export_columns(selected: Iterable[str] | None = None):
    if selected:
        selected_set = set(selected)
        return [col for col in EXPORT_COLUMNS if col[0] in selected_set]
    return EXPORT_COLUMNS[:]


def calculate_next_run(frequency: str, base_time=None):
    base = base_time or timezone.now()
    if frequency == ScheduledExport.Frequency.DAILY:
        return base + timedelta(days=1)
    if frequency == ScheduledExport.Frequency.WEEKLY:
        return base + timedelta(days=7)
    return base + timedelta(days=30)


def _safe_sheet_title(title: str | None) -> str:
    """Return an Excel-compatible worksheet title."""
    cleaned = "".join(ch if ch not in r'[]:*?/\\' else " " for ch in (title or "Export")).strip()
    return (cleaned or "Export")[:31]


def _item_image_url(item: InventoryItem) -> str:
    if not item.image:
        return "–"
    url = item.image.url
    base_url = getattr(settings, "INVENTORY_BASE_URL", "") or ""
    if base_url and url.startswith("/"):
        return f"{base_url.rstrip('/')}{url}"
    return url


def _write_xlsx(
    path: str,
    column_headers: Iterable[object],
    data_rows: Iterable[Iterable[object]],
    sheet_title: str = "Export",
) -> None:
    """Write a formatted .xlsx workbook using openpyxl."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_title(sheet_title)
    worksheet.freeze_panes = "A2"

    headers = list(column_headers)
    rows = [list(row) for row in data_rows]

    header_fill = PatternFill("solid", fgColor="2C3E50")
    odd_fill = PatternFill("solid", fgColor="FFFFFF")
    even_fill = PatternFill("solid", fgColor="F5F6FA")
    thin_side = Side(style="thin", color="DADCE0")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_font = Font(bold=True, color="FFFFFF")
    default_alignment = Alignment(vertical="top", wrap_text=True)
    number_alignment = Alignment(horizontal="right", vertical="top")

    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    worksheet.auto_filter.ref = worksheet.dimensions

    max_widths = [len(str(header or "")) for header in headers]
    for row_idx, row in enumerate(worksheet.iter_rows(), start=1):
        fill = header_fill if row_idx == 1 else (even_fill if row_idx % 2 == 0 else odd_fill)
        for col_idx, cell in enumerate(row, start=1):
            cell.border = border
            cell.fill = fill
            if row_idx == 1:
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.alignment = number_alignment
            else:
                cell.alignment = default_alignment

            if col_idx <= len(headers) and headers[col_idx - 1] == "Bild" and isinstance(cell.value, str):
                if cell.value and cell.value != "–":
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")

            if col_idx > len(max_widths):
                max_widths.append(0)
            text_length = len(str(cell.value or ""))
            max_widths[col_idx - 1] = min(max(max_widths[col_idx - 1], text_length), 40)

    for col_idx, width in enumerate(max_widths, start=1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 42)

    workbook.save(path)


def export_overview_to_file(
    overview: Overview,
    export_format: str | None = None,
    columns: Iterable[str] | None = None,
) -> str:
    export_format = (export_format or "csv").lower()
    if export_format == "excel":
        export_format = "xlsx"

    supported_formats = {"csv", "xlsx", "pdf"}
    if export_format not in supported_formats:
        raise ValueError(f"Unsupported export format: {export_format}")

    selected_columns = get_export_columns(columns)
    extension = export_format
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"overview_{overview.slug}_{timestamp}.{extension}"

    export_dir = os.path.join(settings.MEDIA_ROOT, "exports")
    os.makedirs(export_dir, exist_ok=True)
    full_path = os.path.join(export_dir, filename)

    if export_format == "pdf":
        from .pdf_export import export_overview_to_pdf

        export_overview_to_pdf(overview, selected_columns, full_path)
        return os.path.join("exports", filename)

    items = (
        InventoryItem.objects.filter(overview=overview)
        .select_related("category", "storage_location", "overview")
        .prefetch_related("application_tags")
        .order_by("name")
    )

    column_headers = [col[1] for col in selected_columns]
    data_rows = ([col[2](item) for col in selected_columns] for item in items)

    if export_format == "csv":
        with open(full_path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(column_headers)
            writer.writerows(data_rows)
    elif export_format == "xlsx":
        _write_xlsx(full_path, column_headers, data_rows, sheet_title=overview.name)

    return os.path.join("exports", filename)
